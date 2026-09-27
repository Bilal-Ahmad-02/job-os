"""One-time import of the owner's Job Application Log format.

Never saves the workbook, evaluates formulas, or follows external links.
"""

import hashlib
import warnings
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.xml import DEFUSEDXML
from sqlalchemy import Engine
from sqlalchemy.orm import Session

from app.models.applications import ImportBatch, ImportedRow
from app.schemas.applications import ApplicationData
from app.services.applications import insert, now

SHEET = "Job Application Log"
HEADERS = (
    "Job Title",
    "Company Name/ Website",
    "Website",
    "Where did I find the application for job?",
    "Need to learn about…!",
    "Resume Sent (Date)",
    "How Sent",
    "References Sent",
    "Job Description/ Keywords",
    "Application Status/ Date",
    "Interview/ Date",
    "Follow Up (How/Date)",
    "How I Heard About This Job and Other Comments",
)
FIELDS = (
    "title",
    "company",
    "website",
    "source",
    "learning",
    "resume_sent",
    "how_sent",
    "references_sent",
    "description",
    "status_notes",
    "interview",
    "follow_up",
    "notes",
)
MAX_FILE = 5 * 1024 * 1024
MAX_EXPANDED = 20 * 1024 * 1024


@dataclass(frozen=True)
class SourceRow:
    number: int
    data: ApplicationData
    original: dict[str, str]
    links: dict[str, str]


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time().isoformat() == "00:00:00"
            else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def parse(raw: bytes) -> list[SourceRow]:
    if len(raw) > MAX_FILE or not DEFUSEDXML:
        raise ValueError("Workbook exceeds limits or XML protection is unavailable.")
    # Bound expansion before openpyxl loads cells/styles in memory. Inspect XML
    # dimensions/counts independently: malicious sheets can lie about dimensions.
    with ZipFile(BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 200 or sum(item.file_size for item in entries) > MAX_EXPANDED:
            raise ValueError("Workbook exceeds import limits.")
        if any(item.filename.endswith("vbaProject.bin") for item in entries):
            raise ValueError("Macro-enabled workbooks are not supported.")
        for item in entries:
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                tree = fromstring(archive.read(item))
                ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                rows = list(tree.iter(ns + "row"))
                cells = list(tree.iter(ns + "c"))
                if len(rows) > 1002 or len(cells) > 15000:
                    raise ValueError("Workbook exceeds row/cell limits.")
                if any(int(row.get("r", "0")) > 1002 for row in rows):
                    raise ValueError("Workbook exceeds row limits.")
    # Excel print headers/footers are not application data. Some templates contain
    # a print-footer format openpyxl cannot interpret; cells remain readable.
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Cannot parse header or footer so it will be ignored",
            category=UserWarning,
            module="openpyxl.worksheet.header_footer",
        )
        workbook = load_workbook(BytesIO(raw), data_only=False, keep_links=False)
    try:
        if workbook.sheetnames != [SHEET]:
            raise ValueError("Expected a single Job Application Log sheet.")
        sheet = workbook[SHEET]
        if sheet.max_row > 1002 or sheet.max_column > 14:
            raise ValueError("Unexpected worksheet dimensions.")
        headers = tuple(
            " ".join(cell_text(sheet.cell(2, col).value).split()) for col in range(2, 15)
        )
        if headers != HEADERS:
            raise ValueError("The spreadsheet columns do not match the supported log format.")
        records = []
        for row in sheet.iter_rows(min_row=3, min_col=2, max_col=14):
            if any(cell.data_type in ("f", "e") for cell in row):
                raise ValueError("Resolve formulas or Excel errors before importing.")
            values = {key: cell_text(cell.value) for key, cell in zip(FIELDS, row, strict=True)}
            links = {}
            for cell in row:
                if cell.comment:
                    raise ValueError("Cell comments require manual review before import.")
                if cell.hyperlink:
                    link = cell.hyperlink
                    target = link.target or ""
                    if link.location:
                        target += "#" + link.location
                    if len(target) > 10000:
                        raise ValueError("A hyperlink exceeds the import limit.")
                    links[cell.coordinate] = target
            if not any(value.strip() for value in values.values()) and not links:
                continue
            records.append(SourceRow(row[0].row, ApplicationData(**values), values, links))
        if not records:
            raise ValueError("The spreadsheet has no application records.")
        return records
    finally:
        workbook.close()


def import_workbook(engine: Engine, path: Path) -> dict:
    with path.open("rb") as stream:
        raw = stream.read(MAX_FILE + 1)
    rows = parse(raw)  # Validate every row before touching persistent records.
    digest = hashlib.sha256(raw).hexdigest()
    with Session(engine.execution_options(oracle_write=True)) as session, session.begin():
        previous = session.get(ImportBatch, digest)
        if previous is not None:
            return {"imported": 0, "already_imported": previous.row_count, "sha256": digest}
        session.add(
            ImportBatch(sha256=digest, filename=path.name, imported_at=now(), row_count=len(rows))
        )
        session.flush()
        for row in rows:
            application = insert(session, str(uuid4()), row.data)
            session.add(
                ImportedRow(
                    batch_id=digest,
                    application_id=application.id,
                    sheet=SHEET,
                    row_number=row.number,
                    original=row.original,
                    links=row.links,
                )
            )
    return {"imported": len(rows), "already_imported": 0, "sha256": digest}
